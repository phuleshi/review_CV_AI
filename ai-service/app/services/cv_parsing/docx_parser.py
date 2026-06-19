"""DOCX text extraction using python-docx.

Pulls text from both paragraphs and tables — many CV templates lay content out
in invisible tables, which a paragraph-only read would silently drop.
"""
import io
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.core.exceptions import CVParsingError
from app.core.logging import get_logger
from app.domain.interfaces.extractor import TextExtractor

logger = get_logger(__name__)


class DOCXParser(TextExtractor):
    suffixes = (".docx",)
    content_types = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    def extract(self, content: bytes) -> str:
        if not content:
            raise CVParsingError("DOCX file is empty")
        try:
            doc = Document(io.BytesIO(content))
        except (PackageNotFoundError, BadZipFile, ValueError, KeyError, OSError) as exc:
            raise CVParsingError(f"Could not read DOCX: {exc}") from exc

        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))

        text = "\n".join(lines).strip()
        if not text:
            raise CVParsingError("No extractable text in DOCX")
        logger.info("Extracted %d chars from DOCX", len(text))
        return text
